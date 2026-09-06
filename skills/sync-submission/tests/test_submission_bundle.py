#!/usr/bin/env python3
"""Synthetic submission-bundle controls; no private or third-party fixture files."""
import contextlib
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
import sync_submission as sync


class BundleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name).resolve()
        self.canonical = self.put("manuscript/manuscript.md", "# Synthetic study\n\nSample size: 30.\n")
        self.put("build/supplement.md", "# Synthetic supplement\n\nSample size: 30.\n")
        self.put("build/final.pdf", b"%PDF-1.4\nsynthetic byte-copy fixture; not a render\n")
        self.put("artifact_manifest.json", json.dumps({"schema_version": 1, "artifacts": ["keep"],
                                                      "submissions": {"other": {"status": "submitted"}}}))
        self.spec = {"schema_version": 1, "artifacts": [
            {"id": "supplement", "role": "supplement", "source": "build/supplement.md",
             "target": "supplement/supplement.md", "derived_from": [self.dep(self.canonical)]},
            {"id": "pdf", "role": "manuscript_pdf", "source": "build/final.pdf",
             "target": "manuscript/final.pdf", "transformation": {"kind": "rendered", "command": ["record-only"]},
             "derived_from": [self.dep(self.canonical)], "rights": {"status": "original"}}]}
        self.directory = self.root / "submission/example"

    def put(self, relative, data):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(data if isinstance(data, bytes) else data.encode())
        return path

    def dep(self, path):
        return {"path": path.relative_to(self.root).as_posix(), "sha256": sync.sha256_file(path)}

    def build(self, spec=None):
        with contextlib.redirect_stdout(io.StringIO()):
            return sync.build(self.root, "example", self.canonical, self.spec if spec is None else spec)

    def audit(self):
        with contextlib.redirect_stdout(io.StringIO()):
            code = sync.audit(self.root, "example", self.canonical)
        return code, sync.load_json(self.root / "qc/submission_sync_example.json")

    def snapshot(self):
        return {p.relative_to(self.root).as_posix(): (p.read_bytes(), p.stat().st_mtime_ns)
                for p in self.root.rglob("*") if p.is_file()}

    def test_copy_preserves_bytes_mtime_sources_and_manifest_fields(self):
        before = self.snapshot()
        self.assertEqual(self.build(), 0)
        for path in ("manuscript/manuscript.md", "build/supplement.md", "build/final.pdf"):
            self.assertEqual(self.snapshot()[path], before[path])
        meta = sync.load_json(self.directory / ".journal_meta.json")
        for row in meta["artifacts"]:
            self.assertEqual((self.root / row["source"]).read_bytes(), (self.directory / row["target"]).read_bytes())
        manifest = sync.load_json(self.root / "artifact_manifest.json")
        self.assertEqual(manifest["artifacts"], ["keep"])
        self.assertEqual(manifest["submissions"]["other"], {"status": "submitted"})
        self.assertEqual(manifest["submissions"]["example"]["artifacts"], meta["artifacts"])

    def test_rerun_is_current_but_not_reviewed(self):
        self.build()
        binding = sync.bundle_binding(self.root, "example")
        self.build()
        self.assertEqual(binding, sync.bundle_binding(self.root, "example"))
        code, report = self.audit()
        self.assertEqual(code, 0)
        self.assertEqual(report["readiness"], "not_assessed")
        self.assertEqual(report["verification"]["status"], "not_run")
        self.assertEqual(report["artifacts"][2]["visual_review"], "not_assessed")

    def test_stale_supplement_not_reblessed_after_main_changes(self):
        self.build()
        self.canonical.write_text("# Synthetic study\n\nSample size: 40.\n")
        before = self.snapshot()
        with self.assertRaisesRegex(ValueError, "dependency"):
            self.build()
        self.assertEqual(before, self.snapshot())
        code, report = self.audit()
        self.assertEqual(code, 1)
        self.assertIn("render_input_changed_or_missing", report["artifacts"][1]["drift"])
        self.assertEqual(self.audit()[0], 1)

    def test_edited_output_is_not_overwritten_or_frozen(self):
        self.build()
        edited = self.put("submission/example/manuscript/final.pdf", b"manually edited")
        before = edited.read_bytes()
        with self.assertRaisesRegex(ValueError, "edited"):
            self.build()
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(sync.freeze(self.root, "example", self.canonical, "submitted"), 1)
        self.assertEqual(edited.read_bytes(), before)

    def test_missing_output_and_source_are_drift(self):
        self.build()
        (self.directory / "supplement/supplement.md").unlink()
        (self.root / "build/final.pdf").unlink()
        code, report = self.audit()
        self.assertEqual(code, 1)
        self.assertIn("output_changed_or_missing", report["artifacts"][1]["drift"])
        self.assertIn("source_changed_or_missing", report["artifacts"][2]["drift"])

    def test_missing_submission_cannot_freeze(self):
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(sync.freeze(self.root, "example", self.canonical, "submitted"), 2)
        self.assertFalse(self.directory.exists())

    def test_frozen_bundle_immutable_and_freeze_idempotent(self):
        self.build()
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(sync.freeze(self.root, "example", self.canonical, "submitted"), 0)
            original = (self.directory / ".journal_meta.json").read_bytes()
            self.assertEqual(sync.freeze(self.root, "example", self.canonical, "submitted"), 0)
        self.assertEqual((self.directory / ".journal_meta.json").read_bytes(), original)
        with self.assertRaisesRegex(ValueError, "immutable"):
            self.build()

    def test_unknown_rights_preserved_no_legal_approval(self):
        self.build()
        row = self.audit()[1]["artifacts"][1]
        self.assertEqual(row["rights"]["status"], "unknown")
        self.assertEqual(row["content_fidelity"], "not_assessed")

    def test_incomplete_documented_rights_rejected(self):
        self.spec["artifacts"][0]["rights"] = {"status": "documented", "license_or_permission": "CC BY 4.0"}
        with self.assertRaisesRegex(ValueError, "Documented rights"):
            self.build()
        self.assertFalse(self.directory.exists())

    def test_rights_attribution_and_modification_notice_roundtrip(self):
        rights = {"status": "documented", "source": "synthetic rights example",
                  "license_or_permission": "synthetic permission record", "attribution": "Synthetic creator",
                  "changes": "Example wording adapted"}
        self.spec["artifacts"][0]["rights"] = rights
        self.build()
        self.assertEqual(self.audit()[1]["artifacts"][1]["rights"], rights)

    def test_hidden_and_unregistered_files_remain_visible(self):
        self.build()
        self.put("submission/example/.hidden.txt", "synthetic")
        code, report = self.audit()
        self.assertEqual(code, 1)
        self.assertEqual(report["unregistered_files"], [".hidden.txt"])
        with self.assertRaisesRegex(ValueError, "Unregistered"):
            self.build()

    def test_paths_reject_traversal_symlinks_and_hidden_inputs(self):
        for value in ("../outside", "/tmp/absolute", ".hidden", "nested/../outside", "C:/absolute", "x\\y"):
            with self.subTest(value=value), self.assertRaises(ValueError):
                sync.safe_path(self.root, value)
        (self.root / "alias").symlink_to(self.root / "build", target_is_directory=True)
        with self.assertRaises(ValueError):
            sync.safe_path(self.root, "alias/final.pdf")

    def test_bad_journal_is_rejected_before_writes(self):
        for journal in ("../bad", "nested/path", ".hidden", ""):
            with self.subTest(journal=journal), self.assertRaises(ValueError):
                sync.journal_root(self.root, journal)

    def test_case_duplicate_ancestor_and_reserved_target_rejected(self):
        for target in ("manuscript/MANUSCRIPT.md", "MANUSCRIPT", ".journal_meta.json"):
            self.spec["artifacts"][0]["target"] = target
            with self.subTest(target=target), self.assertRaises(ValueError):
                self.build()

    def test_hardlink_source_output_alias_rejected(self):
        self.directory.mkdir(parents=True)
        dest = self.directory / "manuscript/manuscript.md"
        dest.parent.mkdir()
        os.link(self.canonical, dest)
        before = self.canonical.read_bytes()
        with self.assertRaisesRegex(ValueError, "alias"):
            self.build()
        self.assertEqual(self.canonical.read_bytes(), before)

    def test_missing_declared_input_leaves_prior_package_unchanged(self):
        self.build()
        self.spec["artifacts"][0]["source"] = "build/missing.md"
        before = self.snapshot()
        with self.assertRaises(ValueError):
            self.build()
        self.assertEqual(before, self.snapshot())

    def test_removed_registered_artifact_rejected(self):
        self.build()
        self.spec["artifacts"].pop()
        with self.assertRaisesRegex(ValueError, "remove"):
            self.build()

    def test_malformed_manifest_not_silently_replaced(self):
        self.put("artifact_manifest.json", "{broken")
        before = self.snapshot()
        with self.assertRaises(ValueError):
            self.build()
        self.assertEqual(before, self.snapshot())

    def test_manifest_failure_rolls_back_package(self):
        self.build()
        before = {p: b for p, (b, _) in self.snapshot().items()}
        original = sync.write_json
        def failing(path, payload):
            if path == self.root / "artifact_manifest.json":
                raise OSError("synthetic disk failure")
            return original(path, payload)
        with patch.object(sync, "write_json", side_effect=failing), self.assertRaises(OSError):
            self.build()
        self.assertEqual(before, {p: b for p, (b, _) in self.snapshot().items()})

    def test_lock_does_not_remove_another_writers_lock(self):
        with sync.mutation_lock(self.root):
            with self.assertRaises(ValueError):
                with sync.mutation_lock(self.root):
                    pass
            self.assertTrue((self.root / ".submission-sync.lock").exists())
        self.assertFalse((self.root / ".submission-sync.lock").exists())

    def test_rendered_inputs_must_be_pinned(self):
        self.spec["artifacts"][1]["derived_from"] = []
        with self.assertRaisesRegex(ValueError, "pinned"):
            self.build()

    def test_source_change_during_copy_keeps_prior_bundle(self):
        self.build()
        before = (self.directory / "manuscript/manuscript.md").read_bytes()
        original = sync.shutil.copy2
        def changing(src, dst):
            result = original(src, dst)
            if src == self.canonical:
                self.canonical.write_text("Changed during copy")
            return result
        with patch.object(sync.shutil, "copy2", side_effect=changing), self.assertRaises(ValueError):
            self.build()
        self.assertEqual((self.directory / "manuscript/manuscript.md").read_bytes(), before)

    def test_copied_hidden_docx_metadata_reaches_existing_check(self):
        from docx import Document
        source = self.root / "build/final.docx"
        document = Document()
        document.add_paragraph("Synthetic metadata control")
        document.save(source)
        with zipfile.ZipFile(source, "a") as archive:
            archive.writestr("docProps/custom.xml", '<Properties><property name="source">'
                             '<value>/Users/testuser/styles/journal.csl</value></property></Properties>')
        self.spec["artifacts"][1].update({"source": "build/final.docx", "target": "manuscript/final.docx"})
        self.build()
        self.assertEqual(source.read_bytes(), (self.directory / "manuscript/final.docx").read_bytes())
        report = self.root / "metadata-check.json"
        proc = subprocess.run([sys.executable, str(SCRIPTS / "check_asset_anonymization.py"),
                               "--dir", str(self.directory), "--out", str(report), "--quiet"], capture_output=True)
        self.assertEqual(proc.returncode, 1, proc.stderr)
        findings = json.loads(report.read_text())["findings"]
        self.assertTrue(any(f["type"] == "docx_embedded_abs_path" for f in findings))

    def test_preflight_binding_stays_stale_on_repeated_audit(self):
        self.build()
        self.put("qc/preflight_gate_report.json", json.dumps({"journal": "example",
            "bundle_binding": sync.bundle_binding(self.root, "example"),
            "bundle_unchanged_during_checks": True, "checks": [{"id": "synthetic", "status": "skipped"}]}))
        self.assertEqual(self.audit()[1]["verification"]["status"], "package_bytes_current")
        self.put("submission/example/manuscript/final.pdf", b"changed")
        for _ in range(2):
            self.assertEqual(self.audit()[1]["verification"]["status"], "stale")

    def test_legacy_preflight_is_unbound_not_verified(self):
        self.build()
        self.put("qc/preflight_gate_report.json", json.dumps({"journal": "example", "submission_safe": True}))
        self.assertEqual(self.audit()[1]["verification"]["status"], "unbound")

    def test_yaml_comments_and_nested_paths_do_not_retarget_canonical(self):
        self.put("project.yaml", 'canonical_manuscript: "manuscript/manuscript.md" # source\nother:\n  canonical_manuscript: ignored.md\n')
        self.assertEqual(sync.resolve_canonical(self.root, None), self.canonical)

    def test_explicit_canonical_must_match_recorded_path(self):
        self.build()
        alternate = self.put("alternate.md", self.canonical.read_bytes())
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(sync.audit(self.root, "example", alternate), 1)

    @unittest.skipUnless(shutil.which("pandoc"), "pandoc required for real DOCX control")
    def test_actual_docx_copy_and_manual_word_edit(self):
        from docx import Document
        docx = self.root / "build/final.docx"
        subprocess.run(["pandoc", str(self.canonical), "-o", str(docx)], check=True)
        self.spec["artifacts"][1].update({"source": "build/final.docx", "target": "manuscript/final.docx"})
        self.build()
        output = self.directory / "manuscript/final.docx"
        self.assertEqual(docx.read_bytes(), output.read_bytes())
        document = Document(output)
        document.add_paragraph("Synthetic manual edit")
        document.save(output)
        self.assertEqual(self.audit()[0], 1)

    def test_cli_build_audit_freeze_contract(self):
        self.put("bundle.json", json.dumps(self.spec))
        command = [sys.executable, str(SCRIPTS / "sync_submission.py"), "build", "--project-root", str(self.root),
                   "--journal", "example", "--bundle-spec", "bundle.json"]
        proc = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        command[2] = "audit"
        proc = subprocess.run(command[:-2], capture_output=True, text=True)
        self.assertEqual(proc.returncode, 0, proc.stderr)

    def test_missing_spec_fails_without_writing_package(self):
        command = [sys.executable, str(SCRIPTS / "sync_submission.py"), "build", "--project-root", str(self.root),
                   "--journal", "example", "--bundle-spec", "missing.json"]
        self.assertEqual(subprocess.run(command, capture_output=True).returncode, 2)
        self.assertFalse(self.directory.exists())


if __name__ == "__main__":
    unittest.main()
