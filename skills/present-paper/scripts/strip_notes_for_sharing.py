#!/usr/bin/env python3
"""Strip all speaker notes from a PPTX file (sharing-ready variant).

Use case: after preparing an academic lecture deck where the speaker notes
contain presenter-only material (foreign-language narrative, pronunciation
hints, self-referential reminders), generate a clean copy whose notes are
empty so the deck can be circulated to the audience or to a senior reviewer
without leaking presenter-only content.

The slide body, figures, layout, and pronunciation/asterisk-bearing strings
(e.g. HLA alleles) are preserved exactly. What is removed is everything a
presenter keeps for themselves:

  * the text of EVERY shape on every notes page — not only the notes
    placeholder. A text box dropped onto the Notes Page view is notes too, and
    clearing only ``notes_text_frame`` let it travel while this printed OK;
  * review comments and their author lists (``ppt/comments/``,
    ``commentAuthors.xml``, ``authors.xml``);
  * the author, last-modified-by and comments fields of ``docProps/core.xml``.

Hidden slides (``show="0"``) are a decision, not a default: they are usually
backup or presenter-only material, so the run stops (exit 2) until you pass
``--drop-hidden`` (remove them from the shared copy) or ``--keep-hidden``.

The result is verified against the raw XML of the written file, not through
python-pptx's view of it. Also re-writes ``docProps/app.xml`` so the
PowerPoint Mac repair dialog is not triggered.

Usage:
    python3 strip_notes_for_sharing.py INPUT.pptx OUTPUT.pptx [--drop-hidden|--keep-hidden]

Exit: 0 clean, 2 hidden slides need a decision or presenter-only text survived.
"""
from __future__ import annotations

import argparse
import posixpath
import re
import shutil
import sys
import zipfile
from pathlib import Path
from typing import List, Set
from xml.etree import ElementTree as ET

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
PKG_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
DC = "{http://purl.org/dc/elements/1.1/}"
CP = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"

# Core properties that name a person or carry a note to self. Title and subject stay: they
# describe the talk, and the audience is meant to see them.
CORE_SCRUB = (f"{DC}creator", f"{CP}lastModifiedBy", f"{DC}description")


def _is_comment_part(name: str) -> bool:
    return (name.startswith("ppt/comments/")
            or name in ("ppt/commentAuthors.xml", "ppt/authors.xml"))


def _is_hidden(slide) -> bool:
    return slide._element.get("show") in ("0", "false")


def _clear_txbody(txbody) -> None:
    """Empty a text body, keeping one empty paragraph (a txBody must have at least one)."""
    paras = txbody.findall(f"{A}p")
    for para in paras:
        txbody.remove(para)
    txbody.append(txbody.makeelement(f"{A}p", {}))


def strip_notes(src: Path, dst: Path, hidden: str = "") -> dict:
    """hidden: "" (refuse if any), "drop" or "keep"."""
    if src.resolve() == dst.resolve():
        raise SystemExit("source and destination must differ")
    prs = Presentation(src)
    hidden_idx = [i for i, sl in enumerate(prs.slides, start=1) if _is_hidden(sl)]
    if hidden_idx and not hidden:
        return {"hidden_undecided": hidden_idx}
    if hidden_idx and hidden == "drop":
        sld_ids = prs.slides._sldIdLst
        for sld_id in list(sld_ids):
            slide = prs.slides.get(int(sld_id.get("id")))
            if slide is not None and _is_hidden(slide):
                prs.part.drop_rel(sld_id.rId)
                sld_ids.remove(sld_id)
    cleared = 0
    for slide in prs.slides:
        if not slide.has_notes_slide:
            continue
        el = slide.notes_slide._element
        had_text = any((t.text or "").strip() for t in el.iter(f"{A}t"))
        # Every text body on the notes page — the notes placeholder, a text box added in Notes
        # Page view, a table cell — not only notes_text_frame.
        for txbody in list(el.iter(f"{P}txBody")) + list(el.iter(f"{A}txBody")):
            _clear_txbody(txbody)
        if had_text:
            cleared += 1
    shutil.copy(src, dst)
    prs.save(dst)
    scrub_package(dst)
    return {"cleared": cleared, "total_slides": len(prs.slides),
            "hidden": hidden_idx, "hidden_action": hidden}


def scrub_package(pptx_path: Path) -> None:
    """Drop comment parts (and every reference to them) and blank the personal core properties.

    Elements are removed from the raw text rather than re-serialised through ElementTree, which
    would rename every namespace prefix in files PowerPoint reads strictly.
    """
    with zipfile.ZipFile(pptx_path, "r") as zin:
        names = zin.namelist()
        dropped: Set[str] = {n for n in names if _is_comment_part(n)}
        files = {}
        for n in names:
            if n in dropped:
                continue
            data = zin.read(n)
            if dropped and n.endswith(".rels"):
                owner_dir = posixpath.dirname(posixpath.dirname(n))
                text = data.decode("utf-8")
                for rel in ET.fromstring(data).iter(PKG_REL):
                    if rel.get("TargetMode") == "External":
                        continue
                    t = rel.get("Target", "")
                    tgt = t.lstrip("/") if t.startswith("/") else posixpath.normpath(
                        posixpath.join(owner_dir, t))
                    if tgt in dropped:
                        rid = re.escape(rel.get("Id", ""))
                        text = re.sub(rf'<Relationship\b[^>]*\bId="{rid}"[^>]*/>', "", text)
                data = text.encode("utf-8")
            elif dropped and n == "[Content_Types].xml":
                text = data.decode("utf-8")
                for part in dropped:
                    text = re.sub(rf'<Override\b[^>]*\bPartName="/{re.escape(part)}"[^>]*/>',
                                  "", text)
                data = text.encode("utf-8")
            elif n == "docProps/core.xml":
                text = data.decode("utf-8")
                for local in ("creator", "lastModifiedBy", "description"):
                    text = re.sub(rf"<((?:\w+:)?{local})\b([^>/]*)>.*?</\1>", r"<\1\2></\1>",
                                  text, flags=re.S)
                data = text.encode("utf-8")
            files[n] = data
    tmp = pptx_path.with_suffix(".scrub.tmp.pptx")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n, data in files.items():
            zout.writestr(n, data)
    shutil.move(str(tmp), str(pptx_path))


def fix_app_xml(pptx_path: Path) -> dict:
    """Re-write docProps/app.xml so Slides/Notes counts match reality.

    python-pptx leaves Slides=0 on save, which PowerPoint Mac flags as a
    repair condition. Match canonical pattern in pptx-mac-compatibility.md §5.
    """
    prs = Presentation(pptx_path)
    n_slides = len(prs.slides)
    n_hidden = sum(1 for sl in prs.slides if _is_hidden(sl))
    with zipfile.ZipFile(pptx_path, "r") as z:
        n_notes = sum(
            1 for x in z.namelist()
            if x.startswith("ppt/notesSlides/notesSlide") and x.endswith(".xml")
        )
    titles = []
    for slide in prs.slides:
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                title = shape.text_frame.text.strip().split("\n")[0][:60]
                break
        titles.append(title or "Untitled")

    def esc(t: str) -> str:
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    title_items = "".join(f"<vt:lpstr>{esc(t)}</vt:lpstr>" for t in titles)
    new_app_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        "<TotalTime>1</TotalTime><Words>0</Words>"
        "<Application>Microsoft Macintosh PowerPoint</Application>"
        "<PresentationFormat>Widescreen</PresentationFormat><Paragraphs>0</Paragraphs>"
        f"<Slides>{n_slides}</Slides><Notes>{n_notes}</Notes>"
        f"<HiddenSlides>{n_hidden}</HiddenSlides><MMClips>0</MMClips><ScaleCrop>false</ScaleCrop>"
        '<HeadingPairs><vt:vector size="4" baseType="variant">'
        "<vt:variant><vt:lpstr>Theme</vt:lpstr></vt:variant>"
        "<vt:variant><vt:i4>1</vt:i4></vt:variant>"
        "<vt:variant><vt:lpstr>Slide Titles</vt:lpstr></vt:variant>"
        f"<vt:variant><vt:i4>{n_slides}</vt:i4></vt:variant>"
        "</vt:vector></HeadingPairs>"
        f'<TitlesOfParts><vt:vector size="{n_slides + 1}" baseType="lpstr">'
        f"<vt:lpstr>Office Theme</vt:lpstr>{title_items}"
        "</vt:vector></TitlesOfParts>"
        "<Manager></Manager><Company></Company><LinksUpToDate>false</LinksUpToDate>"
        "<SharedDoc>false</SharedDoc><HyperlinkBase></HyperlinkBase>"
        "<HyperlinksChanged>false</HyperlinksChanged><AppVersion>14.0000</AppVersion>"
        "</Properties>"
    )
    tmp = pptx_path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(pptx_path, "r") as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            zout.writestr(item, new_app_xml if item == "docProps/app.xml" else zin.read(item))
    shutil.move(tmp, pptx_path)
    return {"slides": n_slides, "notes": n_notes}


def verify(pptx_path: Path, keep_hidden: bool = False) -> List[str]:
    """What presenter-only material is still in the written file, read from its raw XML.

    Read through the ZIP, not python-pptx: the earlier check asked python-pptx for
    notes_text_frame — the same object the strip had cleared — so it could only ever agree.
    """
    problems: List[str] = []
    with zipfile.ZipFile(pptx_path) as z:
        for n in z.namelist():
            if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", n):
                left = "".join(t.text or "" for t in ET.fromstring(z.read(n)).iter(f"{A}t"))
                if left.strip():
                    problems.append(f"{n}: {len(left.strip())} chars of notes text remain")
            elif _is_comment_part(n):
                problems.append(f"{n}: review comments remain")
            elif not keep_hidden and re.fullmatch(r"ppt/slides/slide\d+\.xml", n):
                if ET.fromstring(z.read(n)).get("show") in ("0", "false"):
                    problems.append(f"{n}: a hidden slide remains")
            elif n == "docProps/core.xml":
                root = ET.fromstring(z.read(n))
                for tag in CORE_SCRUB:
                    for el in root.iter(tag):
                        if (el.text or "").strip():
                            problems.append(f"{n}: {tag.split('}')[-1]} still set")
    return problems


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("dst", type=Path)
    ap.add_argument("--no-app-xml-fix", action="store_true",
                    help="skip the app.xml rewrite (PowerPoint Mac may show repair dialog)")
    hid = ap.add_mutually_exclusive_group()
    hid.add_argument("--drop-hidden", action="store_true",
                     help="remove hidden slides (show=0) from the shared copy")
    hid.add_argument("--keep-hidden", action="store_true",
                     help="keep hidden slides in the shared copy (they stay hidden)")
    args = ap.parse_args()

    if not args.src.exists():
        print(f"source not found: {args.src}", file=sys.stderr)
        sys.exit(1)

    action = "drop" if args.drop_hidden else "keep" if args.keep_hidden else ""
    info = strip_notes(args.src, args.dst, hidden=action)
    if "hidden_undecided" in info:
        idx = ", ".join(str(i) for i in info["hidden_undecided"])
        print(f"STOP: slide(s) {idx} are hidden. A hidden slide is usually backup or presenter-only "
              "material, and it travels in the file. Re-run with --drop-hidden to remove them "
              "from the shared copy, or --keep-hidden to share them. Nothing was written.",
              file=sys.stderr)
        sys.exit(2)
    print(f"cleared notes on {info['cleared']} / {info['total_slides']} slides")
    if info["hidden"]:
        verb = "removed" if action == "drop" else "kept (still hidden)"
        print(f"hidden slides {verb}: {', '.join(str(i) for i in info['hidden'])}")

    if not args.no_app_xml_fix:
        xml_info = fix_app_xml(args.dst)
        print(f"app.xml patched: Slides={xml_info['slides']}, Notes={xml_info['notes']}")

    leftover = verify(args.dst, keep_hidden=(action == "keep"))
    if leftover:
        print("WARNING: presenter-only material still in the shared copy:", file=sys.stderr)
        for line in leftover:
            print(f"  - {line}", file=sys.stderr)
        sys.exit(2)
    print(f"OK: {args.dst} (notes pages, comments and author fields cleared)")


if __name__ == "__main__":
    main()
