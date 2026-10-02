#!/usr/bin/env python3
"""The checks can only judge what the reader hands them: tables, and slides in the order shown.

Two ways the shared deck reader (check_slide_tells.read_deck, used by check_deck_budget and, for
slide order, check_text_overflow) dropped or misplaced content before any check ran:

  TABLES   A graphicFrame carries its geometry as <p:xfrm>, not <a:xfrm>. The reader looked only
           for <a:xfrm>, so every table was thrown away: a slide holding an 8x5 table of 7-pt text
           passed the density ceiling and the type floor with "OK".
  ORDER    Slides were sorted by the number in slideN.xml, which is a part name, not the order the
           deck is shown in (presentation.xml <p:sldIdLst>). A "Backup" divider moved to the end
           of a 20-slide deck still stopped the clock at slide 3. Notes were paired by the same
           number, so notes created for slide 3 only (notesSlide1.xml) were read as slide 1's.

POSITIVE cases must be flagged; NEGATIVE controls must stay clean.
Skips cleanly (exit 0) without python-pptx. Network-free.

    python3 skills/present-paper/tests/test_deck_reading.py
"""
import re
import sys
import tempfile
import zipfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
except ImportError:
    print("python-pptx not installed — SKIP (compile-only)")
    sys.exit(0)

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import check_deck_budget as budget  # noqa: E402
import check_slide_tells as tells  # noqa: E402
import check_text_overflow as overflow  # noqa: E402

TMP = Path(tempfile.mkdtemp())
fails = []


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    return prs


def textbox(slide, text, y, pt=28):
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(y), Inches(11), Inches(1))
    tb.text_frame.text = text
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(pt)
    return tb


def table_deck(path, rows, cols, cell_text, pt):
    prs = new_deck()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(s, "Ablation halved local recurrence", 0.4, 32)
    tbl = s.shapes.add_table(rows, cols, Inches(0.7), Inches(1.6), Inches(11.5), Inches(5)).table
    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.text = cell_text
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(pt)
    prs.save(str(path))
    return path


def verdicts(path, minutes=10):
    return sorted({f.verdict for f in budget.audit(path, "conference_oral", minutes)})


# --- TABLES ----------------------------------------------------------------------------------
dense = table_deck(TMP / "dense_table.pptx", 8, 5, "alpha beta gamma delta epsilon", 7)
v = verdicts(dense)
for want in ("SLIDE_TOO_DENSE", "TYPE_TOO_SMALL"):
    if want not in v:
        fails.append(f"an 8x5 table of 7-pt text escaped {want} (got {v})")

light = table_deck(TMP / "light_table.pptx", 3, 2, "12% vs 26%", 24)
v = verdicts(light)
if v:
    fails.append(f"a 3x2 table of 24-pt numbers must stay clean, got {v}")
if tells.audit(light):
    fails.append(f"check_slide_tells fired on a plain table slide: "
                 f"{[f.verdict for f in tells.audit(light)]}")

# A table frame whose <p:xfrm> is gone cannot be measured — that is reported, not passed.
broken = TMP / "no_xfrm_table.pptx"
with zipfile.ZipFile(dense) as zin, zipfile.ZipFile(broken, "w", zipfile.ZIP_DEFLATED) as zout:
    for n in zin.namelist():
        data = zin.read(n)
        if n == "ppt/slides/slide1.xml":
            data = re.sub(rb"<p:xfrm>.*?</p:xfrm>", b"", data, flags=re.S)
        zout.writestr(n, data)
if "ZERO_AREA_TEXT" not in verdicts(broken):
    fails.append("a table with no <p:xfrm> was silently skipped instead of reported")
if "ZERO_AREA_TEXT" in verdicts(dense):
    fails.append("ZERO_AREA_TEXT fired on a table whose <p:xfrm> is written")


# --- ORDER -----------------------------------------------------------------------------------
def move_slide_to_end(src, dst, file_index):
    """Move slide<file_index>.xml's <p:sldId> to the end of <p:sldIdLst>, as an editor would,
    leaving every part name unchanged."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        rels = zin.read("ppt/_rels/presentation.xml.rels").decode()
        rid = re.search(rf'Id="(rId\d+)"[^>]*Target="slides/slide{file_index}\.xml"', rels) \
            or re.search(rf'Target="slides/slide{file_index}\.xml"[^>]*Id="(rId\d+)"', rels)
        for n in zin.namelist():
            data = zin.read(n)
            if n == "ppt/presentation.xml":
                text = data.decode()
                ids = re.findall(r"<p:sldId [^>]*/>", text)
                mine = [x for x in ids if f'r:id="{rid.group(1)}"' in x][0]
                body = "".join(x for x in ids if x is not mine) + mine
                text = re.sub(r"(<p:sldIdLst>).*?(</p:sldIdLst>)", rf"\g<1>{body}\g<2>", text,
                              flags=re.S)
                data = text.encode()
            zout.writestr(n, data)
    return dst


heads = [f"Finding {i} changed what we do" for i in range(20)]
heads[2] = "Backup"
prs = new_deck()
for h in heads:
    s = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(s, h, 0.5)
    textbox(s, "eight words of supporting evidence on this slide", 2.5, 24)
in_file_order = TMP / "backup_at_3.pptx"
prs.save(str(in_file_order))
moved = move_slide_to_end(in_file_order, TMP / "backup_moved.pptx", 3)

slides, _n, _w, _h = budget.read_deck(moved)
if budget.find_backup_boundary(slides) != 19:
    fails.append(f"a Backup divider moved to the end must be read as slide 20, got index "
                 f"{budget.find_backup_boundary(slides)}")
if "DECK_OVER_BUDGET" not in verdicts(moved):
    fails.append("19 presented slides before a moved Backup divider passed a 10-min oral")
# control: the same file left in its own order really does have backup at slide 3
if "DECK_OVER_BUDGET" in verdicts(in_file_order):
    fails.append("control deck (Backup really at slide 3) must not be over budget")

# check_text_overflow pairs page N with slide N: N must be the presentation position.
paras = overflow.deck_paragraphs(moved)
if not paras[20] or paras[20][0] != "Backup":
    fails.append(f"check_text_overflow read slide 20 as {paras[20][:1]}, not the moved Backup")

# Notes are found through each slide's own relationship, not by file number.
prs = new_deck()
for i in range(3):
    textbox(prs.slides.add_slide(prs.slide_layouts[6]), f"Slide {i + 1} headline here", 0.5)
prs.slides[2].notes_slide.notes_text_frame.text = "NOTE FOR SLIDE THREE"
notes_deck = TMP / "notes_on_three.pptx"
prs.save(str(notes_deck))
_s, notes, _w, _h = tells.read_deck(notes_deck)
if notes != ["", "", "NOTE FOR SLIDE THREE"]:
    fails.append(f"notes created only for slide 3 were attributed as {notes}")

# A slide id that resolves to no part is an unreadable deck, not a shorter one.
dangling = TMP / "dangling.pptx"
with zipfile.ZipFile(notes_deck) as zin, zipfile.ZipFile(dangling, "w") as zout:
    for n in zin.namelist():
        if n != "ppt/slides/slide2.xml":
            zout.writestr(n, zin.read(n))
try:
    tells.read_deck(dangling)
    fails.append("a deck whose listed slide part is missing was read without complaint")
except KeyError:
    pass

if fails:
    print("FAIL — deck reading")
    for f in fails:
        print("  ·", f)
    sys.exit(1)
print("PASS — tables are measured, slides and notes follow presentation order")
