#!/usr/bin/env python3
"""Exercise every shipped slide builder with synthetic content and existing checks.

Use --write-demo PATH to retain the same deck for a separate renderer-based review.
The tests inspect PPTX structure; they do not certify how PowerPoint renders it.
"""
from pathlib import Path
import sys
import tempfile
import unittest

from PIL import Image
from pptx.util import Pt

SKILL = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(SKILL / 'templates'), str(SKILL / 'scripts')]
import build_pptx_nature_lancet as b
import check_deck_budget as budget

ACADEMIC = ('conference_oral', 'critique', 'case_anchored', 'didactic', 'defence')


def build_demo(path):
    """All seven builders, including a real image slot, short labels and nested lists."""
    prs = b.new_presentation()
    b.add_title_slide(prs, eyebrow='SYNTHETIC EXAMPLE', title='Document processing',
        subtitle='A demonstration of the slide template', meta_top='Software test',
        meta_bottom='Example presenter', notes='Preserve **source text** and *formatting*.')
    b.add_toc_slide(prs, subtitle='Workflow overview', sections=[
        ('01', 'Inputs', 'Read the source files', '5 min'),
        ('02', 'Outputs', 'Check the final document', '5 min')])
    b.add_section_divider(prs, num='01', title='Inputs', subtitle='Text and figures', time_min=5)
    b.add_transition_slide(prs, question='Which file is current?')
    b.add_content_slide(prs, title='Inputs remain unchanged', subtitle='Source preservation',
        bullets=['Keep the **original file**', '  Write outputs to a separate folder',
                 'Check the final artifact'])
    with tempfile.TemporaryDirectory() as td:
        slot = Path(td) / 'slot.png'
        # An inert image-slot fixture, not an illustrative visual asset.
        Image.new('RGB', (120, 160), '#e0e0e0').save(slot)
        b.add_content_slide(prs, title='Output review', bullets=['Read the text', 'Inspect the figure'],
            figure_path=slot, fig_caption='Synthetic image slot', footnote='Synthetic source')
    b.add_glossary_slide(prs, tier1=[('SRC', 'Source file'), ('OUT', 'Output file'),
        ('LOG', 'Run record'), ('VER', 'File version')],
        tier2=[(f'T{i}', f'Example term {i}') for i in range(8)])
    b.add_closing_slide(prs, bullets=['Preserve the source', 'Review the output'],
        contact='Documentation team')
    prs.save(path)
    b.fix_app_xml(path)
    return prs


class BuilderCompatibility(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.deck = Path(self.temp.name) / 'deck.pptx'
        self.prs = build_demo(self.deck)

    def test_every_builder_clears_academic_budget(self):
        self.assertEqual(len(self.prs.slides), 8)
        for archetype in ACADEMIC:
            with self.subTest(archetype=archetype):
                self.assertEqual(budget.audit(self.deck, archetype, 20), [])

    def test_large_room_profiles_are_not_silently_certified(self):
        for archetype in ('keynote', 'lay_talk', 'decision_brief'):
            findings = budget.audit(self.deck, archetype, 20)
            self.assertIn('TYPE_TOO_SMALL', {f.verdict for f in findings})

    def test_regressing_glossary_type_is_detected(self):
        for shape in self.prs.slides[6].shapes:
            if shape.has_text_frame and shape.text.startswith('SRC'):
                for p in shape.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(12)
        self.prs.save(self.deck)
        self.assertIn('TYPE_TOO_SMALL', {f.verdict for f in budget.audit(self.deck, 'conference_oral', 20)})

    def test_native_bullets_preserve_content_and_hanging_indent(self):
        for slide_no, first in [(4, 'Keep the original file'), (7, 'Preserve the source')]:
            shape = next(s for s in self.prs.slides[slide_no].shapes
                         if s.has_text_frame and s.text.startswith(first))
            for p in shape.text_frame.paragraphs:
                self.assertFalse(p.text.startswith(('▪', '—', '–')))
                props = p._p.get_or_add_pPr()
                self.assertIsNotNone(props.find(f'{{{b.A_NS}}}buChar'))
                margin, indent = int(props.get('marL')), int(props.get('indent'))
                self.assertLess(indent, 0)
                self.assertGreater(margin + indent, 0)
                self.assertTrue(all(r.font.size.pt >= 20 for r in p.runs))
            if slide_no == 4:
                self.assertEqual(shape.text_frame.paragraphs[1].level, 1)
                self.assertTrue(any(r.font.bold and r.text == 'original file'
                                    for r in shape.text_frame.paragraphs[0].runs))

    def test_fonts_can_change_without_rewriting_text_or_formatting(self):
        def frames():
            for slide in self.prs.slides:
                yield from (s.text_frame for s in slide.shapes if s.has_text_frame)
                if slide.has_notes_slide:
                    yield slide.notes_slide.notes_text_frame
        def snapshot():
            return [(r.text, r.font.size, r.font.bold, r.font.italic)
                    for tf in frames() for p in tf.paragraphs for r in p.runs]
        before = snapshot()
        b.apply_fonts(self.prs, en='Example Sans', ko='Example CJK')
        self.assertEqual(snapshot(), before)
        for tf in frames():
            for p in tf.paragraphs:
                for r in p.runs:
                    self.assertEqual(r.font.name, 'Example Sans')
                    self.assertEqual(r._r.get_or_add_rPr().find(f'{{{b.A_NS}}}ea').get('typeface'), 'Example CJK')
                for font in p._p.iter(f'{{{b.A_NS}}}buFont'):
                    self.assertEqual(font.get('typeface'), 'Example Sans')
        with self.assertRaises(ValueError):
            b.apply_fonts(self.prs, en='', ko='Example CJK')

    def test_maximum_outline_and_glossary_stay_on_canvas(self):
        prs = b.new_presentation()
        b.add_toc_slide(prs, sections=[(str(i), 'Section', 'Short summary', '5 min') for i in range(5)])
        b.add_glossary_slide(prs, tier1=[(f'A{i}', 'Main concept') for i in range(7)],
                             tier2=[(f'B{i}', 'Short definition') for i in range(12)])
        for slide in prs.slides:
            for shape in slide.shapes:
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, prs.slide_width)
                self.assertLessEqual(shape.top + shape.height, prs.slide_height)
        prs.save(self.deck)
        # Capacity limits are not a promise of acceptable information density.
        findings = budget.audit(self.deck, 'conference_oral', 20)
        self.assertEqual({f.verdict for f in findings}, {'SLIDE_TOO_DENSE'})

    def test_excess_entries_fail_before_creating_a_partial_slide(self):
        before = len(self.prs.slides)
        for call in [lambda: b.add_toc_slide(self.prs, sections=[('1', 'Section', '', '')]*6),
                     lambda: b.add_glossary_slide(self.prs, tier1=[('A', 'Term')]*8),
                     lambda: b.add_glossary_slide(self.prs, tier2=[('A', 'Term')]*13)]:
            with self.assertRaisesRegex(ValueError, 'split'):
                call()
            self.assertEqual(len(self.prs.slides), before)


if __name__ == '__main__':
    if len(sys.argv) == 3 and sys.argv[1] == '--write-demo':
        build_demo(Path(sys.argv[2]))
    else:
        unittest.main()
