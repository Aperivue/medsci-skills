#!/usr/bin/env python3
"""Reproducible test for inject_speaker_notes.py inline-markdown rendering.

Builds a 1-slide PPTX, injects a note containing **bold** / *italic*, and asserts the
asterisks are gone and the runs are styled. --no-markdown must keep the text verbatim.

Also reads the SAVED notes XML for size: every run must carry an explicit ``sz`` (18 pt by
default, ``--font-pt`` otherwise) — a run without one inherits the notes master's 12 pt — and a
blank line must be an empty paragraph with ``<a:endParaRPr sz>``, not an empty run.
Skips cleanly (exit 0) when python-pptx is not installed. Network-free.

    python3 skills/present-paper/tests/test_speaker_notes_markdown.py
"""
import importlib.util
import sys
import tempfile
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    print("python-pptx not installed — SKIP (compile-only)")
    sys.exit(0)

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "inject_speaker_notes.py"
spec = importlib.util.spec_from_file_location("isn", SCRIPT)
m = importlib.util.module_from_spec(spec)
spec.loader.exec_module(m)

tmp = Path(tempfile.mkdtemp())
prs = Presentation()
prs.slides.add_slide(prs.slide_layouts[6])
src = tmp / "in.pptx"
prs.save(str(src))

m.notes = {1: "Plain **bold** and *italic* here.\nSecond line."}
out = tmp / "in_notes.pptx"
m.inject_notes(str(src), str(out), markdown=True)
tf = Presentation(str(out)).slides[0].notes_slide.notes_text_frame
runs = [(r.text, r.font.bold, r.font.italic) for p in tf.paragraphs for r in p.runs]

out2 = tmp / "legacy.pptx"
m.inject_notes(str(src), str(out2), markdown=False)
tf2 = Presentation(str(out2)).slides[0].notes_slide.notes_text_frame

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def note_body(path):
    """The saved notes body, reloaded from disk: what a renderer will read."""
    body = Presentation(str(path)).slides[0].notes_slide.notes_text_frame._txBody
    paras = body.findall(f"{A}p")
    sizes = []
    for r in body.iter(f"{A}r"):
        rpr = r.find(f"{A}rPr")
        sizes.append(rpr.get("sz") if rpr is not None else None)
    return paras, sizes


def blank_is_paragraph(paras, sz):
    """The middle line is blank: no run, and an endParaRPr carrying the size."""
    if len(paras) != 3:
        return False
    mid = paras[1]
    end = mid.find(f"{A}endParaRPr")
    return mid.find(f"{A}r") is None and end is not None and end.get("sz") == sz


m.notes = {1: "First line with **bold**.\n\nThird line."}
sized = tmp / "sized.pptx"
m.inject_notes(str(src), str(sized), markdown=True)
paras_md, sizes_md = note_body(sized)

sized_plain = tmp / "sized_plain.pptx"
m.inject_notes(str(src), str(sized_plain), markdown=False)
paras_pl, sizes_pl = note_body(sized_plain)

# The CLI flag, through argparse, the way a user sets it.
cli_out = tmp / "cli.pptx"
argv = sys.argv
sys.argv = ["inject_speaker_notes.py", str(src), "-o", str(cli_out), "--font-pt", "14"]
try:
    m.main()
    _, sizes_cli = note_body(cli_out)
except SystemExit:
    sizes_cli = []
finally:
    sys.argv = argv

checks = [
    ("no literal '**' in rendered notes", "**" not in tf.text),
    ("no literal '*' in rendered notes", "*" not in tf.text),
    ("a bold run 'bold' exists", any(t == "bold" and b for t, b, i in runs)),
    ("an italic run 'italic' exists", any(t == "italic" and i for t, b, i in runs)),
    ("line structure preserved (>=2 paragraphs)", len(tf.paragraphs) >= 2),
    ("--no-markdown keeps '**bold**' verbatim", "**bold**" in tf2.text),
    ("every note run is 18 pt by default (not the master's 12 pt)",
     bool(sizes_md) and all(sz == "1800" for sz in sizes_md)),
    ("a blank line is an empty paragraph with endParaRPr sz=1800, not an empty run",
     blank_is_paragraph(paras_md, "1800")),
    ("--no-markdown: every run 18 pt", bool(sizes_pl) and all(sz == "1800" for sz in sizes_pl)),
    ("--no-markdown: blank line is an empty paragraph with endParaRPr sz=1800",
     blank_is_paragraph(paras_pl, "1800")),
    ("--font-pt 14 sets every run to 14 pt",
     bool(sizes_cli) and all(sz == "1400" for sz in sizes_cli)),
]
fail = 0
for name, ok in checks:
    print(("ok   " if ok else "FAIL ") + name)
    fail += 0 if ok else 1
print("ALL SPEAKER-NOTES MARKDOWN TESTS PASSED" if not fail else f"{fail} FAILED")
sys.exit(1 if fail else 0)
