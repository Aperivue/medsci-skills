#!/usr/bin/env python3
"""The sharing copy must not carry what the presenter kept for themselves.

strip_notes_for_sharing.py used to clear `notes_text_frame` and then verify `notes_text_frame` —
the same object, so the check could only agree with the strip. A text box added on the Notes Page
view, review comments, the author field and hidden backup slides all travelled in the "OK" copy.

  POSITIVE  a deck with every one of those: the shared copy holds none of them, read from the raw
            ZIP (not through python-pptx), and a deck with a hidden slide is refused until the
            user decides (--drop-hidden / --keep-hidden).
  NEGATIVE  a plain deck with notes only: exit 0, slide bodies byte-for-byte the same text, and
            the slide count unchanged.

Skips cleanly (exit 0) without python-pptx. Network-free.

    python3 skills/present-paper/tests/test_strip_notes_for_sharing.py
"""
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches
except ImportError:
    print("python-pptx not installed — SKIP (compile-only)")
    sys.exit(0)

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "strip_notes_for_sharing.py"
TMP = Path(tempfile.mkdtemp())
fails = []

SECRETS = ("PRIVATE-NOTEBOX", "PLACEHOLDER-NOTE", "HIDDEN-BACKUP", "AUTHOR-NAME",
           "DRAFT-COMMENT", "REVIEW-COMMENT-TEXT")

COMMENT_AUTHORS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<p:cmAuthorLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:cmAuthor id="0" name="Reviewer" initials="R" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>')
COMMENTS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
    '<p:cm authorId="0" dt="2026-01-01T00:00:00.000" idx="1"><p:pos x="10" y="10"/>'
    '<p:text>REVIEW-COMMENT-TEXT</p:text></p:cm></p:cmLst>')


def add_comments(path: Path) -> None:
    """Attach a legacy comment part to slide 1 and a commentAuthors part to the presentation."""
    tmp = path.with_suffix(".c.pptx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            data = zin.read(n).decode("utf-8") if n.endswith((".xml", ".rels")) else zin.read(n)
            if n == "[Content_Types].xml":
                data = data.replace("</Types>", (
                    '<Override PartName="/ppt/comments/comment1.xml" ContentType="application/'
                    'vnd.openxmlformats-officedocument.presentationml.comments+xml"/>'
                    '<Override PartName="/ppt/commentAuthors.xml" ContentType="application/'
                    'vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/>'
                    "</Types>"))
            elif n == "ppt/slides/_rels/slide1.xml.rels":
                data = data.replace("</Relationships>", (
                    '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/comments" Target="../comments/comment1.xml"/>'
                    "</Relationships>"))
            elif n == "ppt/_rels/presentation.xml.rels":
                data = data.replace("</Relationships>", (
                    '<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/'
                    'officeDocument/2006/relationships/commentAuthors" Target="commentAuthors.xml"/>'
                    "</Relationships>"))
            zout.writestr(n, data)
        zout.writestr("ppt/comments/comment1.xml", COMMENTS)
        zout.writestr("ppt/commentAuthors.xml", COMMENT_AUTHORS)
    tmp.replace(path)


def build(path: Path, leaky: bool) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Public body"
    s.notes_slide.notes_text_frame.text = "PLACEHOLDER-NOTE narrative"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "Second body"
    if leaky:
        box = s.shapes.add_textbox(Inches(1), Inches(3), Inches(5), Inches(1))
        box.text_frame.text = "PRIVATE-NOTEBOX: the chair will ask"
        el = box._element
        el.getparent().remove(el)
        s.notes_slide.shapes._spTree.append(el)  # a text box on the Notes Page view
        hid = prs.slides.add_slide(prs.slide_layouts[6])
        hid._element.set("show", "0")
        hid.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = \
            "HIDDEN-BACKUP fallback"
        prs.core_properties.author = "AUTHOR-NAME"
        prs.core_properties.comments = "DRAFT-COMMENT"
    prs.save(str(path))
    if leaky:
        add_comments(path)


def run(*args):
    return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)],
                          capture_output=True, text=True)


def raw_hits(path: Path):
    hits = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            text = z.read(n).decode("utf-8", "ignore")
            hits += [f"{s} in {n}" for s in SECRETS if s in text]
    return hits


# --- POSITIVE: everything presenter-only is gone ----------------------------------------------
leaky = TMP / "leaky.pptx"
build(leaky, leaky=True)
# sanity: the fixture really carries every secret (otherwise the test proves nothing)
if len({h.split(" in ")[0] for h in raw_hits(leaky)}) != len(SECRETS):
    fails.append(f"fixture is missing a planted secret: {raw_hits(leaky)}")

r = run(leaky, TMP / "undecided.pptx")
if r.returncode != 2 or (TMP / "undecided.pptx").exists():
    fails.append(f"a deck with a hidden slide must stop (exit 2, nothing written); got "
                 f"{r.returncode}: {r.stdout}{r.stderr}")

dropped = TMP / "dropped.pptx"
r = run(leaky, dropped, "--drop-hidden")
if r.returncode != 0:
    fails.append(f"--drop-hidden should succeed; got {r.returncode}: {r.stdout}{r.stderr}")
else:
    hits = raw_hits(dropped)
    if hits:
        fails.append(f"presenter-only material survived --drop-hidden: {hits}")
    p = Presentation(str(dropped))
    if len(p.slides) != 2:
        fails.append(f"--drop-hidden should leave 2 slides, left {len(p.slides)}")
    with zipfile.ZipFile(dropped) as z:
        if any(n.startswith("ppt/comments/") or n == "ppt/commentAuthors.xml" for n in z.namelist()):
            fails.append("comment parts survived")
        for n in z.namelist():
            if n.endswith(".rels") and "comment" in z.read(n).decode():
                fails.append(f"a relationship still points at a removed comment part: {n}")
        if "comment" in z.read("[Content_Types].xml").decode():
            fails.append("[Content_Types].xml still declares a removed comment part")
        app = z.read("docProps/app.xml").decode()
        if "<HiddenSlides>0</HiddenSlides>" not in app or "<Slides>2</Slides>" not in app:
            fails.append(f"app.xml counts wrong after --drop-hidden: {app[:300]}")

kept = TMP / "kept.pptx"
r = run(leaky, kept, "--keep-hidden")
if r.returncode != 0:
    fails.append(f"--keep-hidden should succeed; got {r.returncode}: {r.stdout}{r.stderr}")
else:
    with zipfile.ZipFile(kept) as z:
        app = z.read("docProps/app.xml").decode()
    m = re.search(r"<HiddenSlides>(\d+)</HiddenSlides>", app)
    if not m or m.group(1) != "1":
        fails.append(f"--keep-hidden: app.xml must report the 1 hidden slide, got {m and m.group(1)}")
    left = [h for h in raw_hits(kept) if not h.startswith("HIDDEN-BACKUP")]
    if left:
        fails.append(f"--keep-hidden: notes/comments/author survived: {left}")

# --- NEGATIVE: a plain deck is stripped of its notes and nothing else --------------------------
plain = TMP / "plain.pptx"
build(plain, leaky=False)
shared = TMP / "plain_shared.pptx"
r = run(plain, shared)
if r.returncode != 0 or "OK:" not in r.stdout:
    fails.append(f"a plain deck must strip cleanly (exit 0, OK); got {r.returncode}: "
                 f"{r.stdout}{r.stderr}")
else:
    a, b = Presentation(str(plain)), Presentation(str(shared))
    body = lambda p: [[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in p.slides]
    if body(a) != body(b):
        fails.append(f"slide bodies changed: {body(a)} -> {body(b)}")
    if raw_hits(shared):
        fails.append(f"notes survived on a plain deck: {raw_hits(shared)}")

if fails:
    print("FAIL — strip_notes_for_sharing")
    for f in fails:
        print("  ·", f)
    sys.exit(1)
print("PASS — notes pages, comments, author fields and hidden slides handled; plain deck intact")
