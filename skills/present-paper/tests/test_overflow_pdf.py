#!/usr/bin/env python3
"""Exercise the real poppler command, without needing an office renderer in CI.

The tiny vector PDFs below are controlled measurements, not exports of the PPTX.
Builder rendering and PowerPoint compatibility still need separate visual review.
"""
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
import check_text_overflow as overflow


def write_pdf(path, *, baseline=400):
    """One page, standard Helvetica, one line at a known PDF baseline."""
    stream = f'BT /F1 20 Tf 72 {baseline} Td (Synthetic text) Tj ET\n'.encode('ascii')
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 960 540] '
        b'/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        f'<< /Length {len(stream)} >>\nstream\n'.encode('ascii') + stream + b'endstream']
    data = bytearray(b'%PDF-1.4\n'); offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data)); data.extend(f'{i} 0 obj\n'.encode('ascii') + obj + b'\nendobj\n')
    xref = len(data)
    data.extend(f'xref\n0 {len(offsets)}\n0000000000 65535 f \n'.encode('ascii'))
    for offset in offsets[1:]:
        data.extend(f'{offset:010d} 00000 n \n'.encode('ascii'))
    data.extend(f'trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode('ascii'))
    path.write_bytes(data)


class ActualPDFMeasurement(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.deck = Path(self.temp.name) / 'deck.pptx'
        self.pdf = Path(self.temp.name) / 'deck.pdf'
        prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        # A background and a foreground text box deliberately overlap. The
        # overlap itself is not clipping and must remain a clean control.
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.8), Inches(1.5), Inches(8), Inches(2))
        card.fill.solid()
        slide.shapes.add_textbox(Inches(1), Inches(1.7), Inches(7), Inches(1)).text = 'Synthetic text'
        prs.save(self.deck)

    def measure(self, baseline):
        write_pdf(self.pdf, baseline=baseline)
        xml = overflow.run_pdftotext(self.pdf)
        self.assertIn('<line ', xml)
        self.assertIn('Synthetic', xml)
        return overflow.audit(self.deck, xml, .1)

    def test_real_pdf_text_inside_filled_background_is_clean(self):
        self.assertEqual(self.measure(400), [])

    def test_real_pdf_line_crossing_card_is_reported(self):
        self.assertEqual({f.verdict for f in self.measure(282)}, {'CARD'})

    def test_real_pdf_line_crossing_bottom_reserve_is_reported(self):
        self.assertEqual({f.verdict for f in self.measure(6)}, {'OFF_SLIDE'})
        result = subprocess.run([sys.executable, str(Path(overflow.__file__)), str(self.deck),
            '--pdf', str(self.pdf), '--strict'], text=True, capture_output=True)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)

    def test_plain_bbox_is_not_misreported_as_a_clean_measurement(self):
        write_pdf(self.pdf)
        xml = subprocess.run(['pdftotext', '-bbox', str(self.pdf), '-'],
            check=True, text=True, capture_output=True).stdout
        self.assertIn('<word ', xml)
        self.assertNotIn('<line ', xml)
        with self.assertRaisesRegex(overflow.CannotMeasure, 'bbox-layout'):
            overflow.audit(self.deck, xml, .1)

    def test_blank_page_is_not_a_parser_error(self):
        self.assertEqual(overflow.parse_bbox('<page width="960" height="540"></page>'), {1: []})

    def test_one_usable_page_does_not_hide_an_incompatible_page(self):
        xml = ('<page width="960" height="540"><line xMin="1" yMin="2" xMax="3" yMax="4">'
               '<word>Readable</word></line></page><page width="960" height="540">'
               '<word xMin="1" yMin="2" xMax="3" yMax="4">Unusable</word></page>')
        with self.assertRaises(overflow.CannotMeasure):
            overflow.parse_bbox(xml)


if __name__ == '__main__':
    unittest.main()
