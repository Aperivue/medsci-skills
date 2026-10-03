#!/usr/bin/env python3
"""md_marker_convert.py — Convert Vancouver-style ``[N]`` / ``[N, M]`` markers in
markdown or .docx into pandoc-style ``[@KEY]`` / ``[@KEY1, @KEY2]`` citekeys
(or back), driven by an explicit N→key mapping.

Two directions:
  --to-keys     ``[N]``  → ``[@KEY]``   (default)
  --to-numbers  ``[@KEY]`` → ``[N]``    (round-trip / debug)

Inputs:
  --input  FILE        path to .md or .docx
  --output FILE        write converted file (extension must match input)
  --map    FILE        N↔key mapping (JSON ``{"1": "ABC123", ...}`` or 2-column CSV ``n,key``)
  --active-ns CSV      optional comma list of N's to convert; markers outside
                       this set are left untouched. Use to stage partial
                       conversion (e.g. sample-only build).

A range inside a marker (``[1-3]``, ``[1–3]``, ``[2, 4-6]``) is expanded to
its members before conversion, so ``[1-3]`` becomes ``[@K1, @K2, @K3]``.

Numbers without a mapping (or outside ``--active-ns``) are left as plain
``[N]`` so a downstream Zotero refresh / hand-edit can finalize them.

Anti-Hallucination:
  - The mapping is the single source of truth. Never invents keys.
  - Every marker left untouched is reported on stderr, with the reason:
    a number with no mapping (UNMAPPED), a number outside ``--active-ns``
    (INACTIVE, the staged-conversion case), or a range that runs backwards
    (MALFORMED).
  - Exit 1 when a marker was left because of an UNMAPPED number or a
    MALFORMED range (or, with ``--to-numbers``, a key not in the map); the
    output is still written. Pass ``--allow-partial`` to accept that partial
    conversion with exit 0. INACTIVE markers alone never fail: leaving them
    is what ``--active-ns`` asks for.
  - The pattern is chosen by direction, never by the file's content:
    ``--to-keys`` converts only ``[N]`` markers and leaves ``[@key]`` text
    alone, so re-running it on a partially converted file is safe.

Origin: generalized from the a per-project ``build_zotero_docx.py`` replacer
(2026-05-01), validated on a 21-reference manuscript.
"""

from __future__ import annotations

import argparse
import csv
import itertools
import json
import re
import sys
from pathlib import Path

# [N], [N, M], [N,M,...], possibly with internal whitespace; any member may be a range N-M
# (hyphen, en dash or em dash). A range used to fall outside the pattern entirely, so "[1-3]" was
# neither converted nor counted nor reported.
_RANGE_SEP = r"\s*[-\u2013\u2014]\s*"
_NUM_ITEM = rf"\d+(?:{_RANGE_SEP}\d+)?"
_NUM_RE = re.compile(rf"\[({_NUM_ITEM}(?:\s*,\s*{_NUM_ITEM})*)\]")
# Pandoc citekey group: [@KEY], [@K1, @K2], [@K1; @K2]
_KEY_RE = re.compile(r"\[(@[A-Za-z0-9_]+(?:\s*[,;]\s*@[A-Za-z0-9_]+)*)\]")


def load_map(path: Path) -> dict[int, str]:
    raw = path.read_text(encoding="utf-8").strip()
    if path.suffix.lower() == ".json" or raw.startswith("{"):
        data = json.loads(raw)
        return {int(k): str(v).strip() for k, v in data.items() if str(v).strip()}
    out: dict[int, str] = {}
    for row in csv.reader(raw.splitlines()):
        if not row or row[0].strip().lower() in {"n", "number", ""}:
            continue
        if len(row) < 2:
            continue
        out[int(row[0].strip())] = row[1].strip()
    return out


def parse_active(spec: str | None, full_keys: set[int]) -> set[int]:
    if not spec:
        return full_keys
    return {int(x.strip()) for x in spec.split(",") if x.strip()}


def expand_marker(body: str) -> list[range] | None:
    """The numbers a marker body names, as one range per item. None for a backwards range.

    Ranges stay lazy: "[1-100000000]" is not materialised. The caller stops at the first number
    the map lacks, so the work done is bounded by the size of the map, not by the range.
    """
    spans: list[range] = []
    for item in body.split(","):
        parts = [x.strip() for x in re.split(_RANGE_SEP, item.strip())]
        if len(parts) == 1:
            n = int(parts[0])
            spans.append(range(n, n + 1))
            continue
        lo, hi = int(parts[0]), int(parts[1])
        if hi < lo:
            return None
        spans.append(range(lo, hi + 1))
    return spans


def make_num_to_key(n_to_key: dict[int, str], active: set[int], staged: bool = False):
    """Replacement function for ``[N]`` markers, plus the record of every marker it left.

    ``left`` maps a reason (UNMAPPED / INACTIVE / MALFORMED) to the markers left for it. A number
    with no mapping used to be dropped before it was recorded whenever it was also outside the
    active set — which, with no --active-ns, is every unmapped number — so it was never reported.
    ``staged`` is True when --active-ns was given: a number outside it is INACTIVE (left on
    purpose) whether or not the map covers it; a number inside it with no mapping is UNMAPPED.
    """
    left: dict[str, list[str]] = {"UNMAPPED": [], "INACTIVE": [], "MALFORMED": []}

    def repl(m: re.Match) -> str:
        spans = expand_marker(m.group(1))
        if spans is None:
            left["MALFORMED"].append(m.group(0))
            return m.group(0)
        if staged:
            # Only an ACTIVE number can be unmapped; test the (finite) active set against the spans.
            unmapped = any(n not in n_to_key and any(n in sp for sp in spans) for n in active)
        else:
            # Stops at the first number the map lacks: at most len(map) + 1 steps.
            unmapped = any(n not in n_to_key for n in itertools.chain.from_iterable(spans))
        if unmapped:
            left["UNMAPPED"].append(m.group(0))
            return m.group(0)
        # Stops at the first number outside the active set: at most len(active) + 1 steps. Past it,
        # every number is active and mapped, so the join is bounded by the map size.
        if not all(n in active for n in itertools.chain.from_iterable(spans)):
            left["INACTIVE"].append(m.group(0))
            return m.group(0)
        return "[" + ", ".join(f"@{n_to_key[n]}" for n in itertools.chain.from_iterable(spans)) + "]"

    return repl, left


def make_key_to_num(n_to_key: dict[int, str]):
    key_to_n = {v: n for n, v in n_to_key.items()}
    unknown: set[str] = set()

    def repl(m: re.Match) -> str:
        keys = [k.strip().lstrip("@") for k in re.split(r"[,;]", m.group(1))]
        nums: list[int] = []
        for k in keys:
            n = key_to_n.get(k)
            if n is None:
                unknown.add(k)
                return m.group(0)
            nums.append(n)
        return "[" + ", ".join(str(n) for n in nums) + "]"

    return repl, unknown


def pattern_for(direction: str) -> re.Pattern:
    """The marker pattern is set by the direction asked for, never guessed from the content.

    Guessing ("[@" present -> key pattern) made a re-run of --to-keys on a partially converted file
    feed "@KEY" to int() and crash, and would have converted nothing on a file that mixed both.
    """
    return _KEY_RE if direction == "to-numbers" else _NUM_RE


def transform_text(text: str, repl, direction: str = "to-keys") -> str:
    return pattern_for(direction).sub(repl, text)


def convert_markdown(src: Path, dst: Path, repl, direction: str = "to-keys") -> int:
    text = src.read_text(encoding="utf-8")
    pattern = pattern_for(direction)
    n = sum(1 for _ in pattern.finditer(text))
    new = pattern.sub(repl, text)
    dst.write_text(new, encoding="utf-8")
    return n


def convert_docx(src: Path, dst: Path, repl, direction: str) -> int:
    try:
        from docx import Document  # type: ignore
    except ImportError:
        sys.exit("ERROR: python-docx is required for .docx input. `pip install python-docx`")
    doc = Document(str(src))
    pattern = pattern_for(direction)
    n = 0

    def process(p):
        nonlocal n
        for run in p.runs:
            if "[" not in run.text:
                continue
            n += sum(1 for _ in pattern.finditer(run.text))
            run.text = pattern.sub(repl, run.text)

    for p in doc.paragraphs:
        process(p)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process(p)
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return n


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--input", required=True, type=Path)
    ap.add_argument("--output", required=True, type=Path)
    ap.add_argument("--map", required=True, type=Path)
    ap.add_argument("--active-ns", default=None,
                    help="Comma-separated N's to convert (default: all in map).")
    direction = ap.add_mutually_exclusive_group()
    direction.add_argument("--to-keys", action="store_true", default=True,
                           help="[N] → [@KEY] (default).")
    direction.add_argument("--to-numbers", action="store_true",
                           help="[@KEY] → [N] (round-trip).")
    ap.add_argument("--allow-partial", action="store_true",
                    help="Exit 0 even when markers were left for an UNMAPPED number, a "
                         "MALFORMED range or an unknown key (they are still reported).")
    args = ap.parse_args()

    if not args.input.exists():
        sys.exit(f"ERROR: input not found: {args.input}")
    if args.input.suffix != args.output.suffix:
        sys.exit("ERROR: --input and --output must share the same extension (.md or .docx).")

    n_to_key = load_map(args.map)
    if not n_to_key:
        sys.exit("ERROR: empty mapping.")

    unknown: set[str] = set()
    left: dict[str, list[str]] = {}
    if args.to_numbers:
        repl, unknown = make_key_to_num(n_to_key)
        direction = "to-numbers"
    else:
        active = parse_active(args.active_ns, set(n_to_key))
        repl, left = make_num_to_key(n_to_key, active, staged=bool(args.active_ns))
        direction = "to-keys"

    if args.input.suffix == ".docx":
        seen = convert_docx(args.input, args.output, repl, direction)
    else:
        seen = convert_markdown(args.input, args.output, repl, direction)

    print(f"[md_marker_convert] direction={direction} markers_seen={seen} → {args.output}",
          file=sys.stderr)
    why = {
        "UNMAPPED": "a number has no entry in the map",
        "MALFORMED": "the range runs backwards",
        "INACTIVE": "outside --active-ns (staged conversion; not a failure)",
    }
    for reason in ("UNMAPPED", "MALFORMED", "INACTIVE"):
        markers = left.get(reason, [])
        if markers:
            label = "NOTE" if reason == "INACTIVE" else "WARNING"
            print(f"[md_marker_convert] {label}: {len(markers)} marker(s) left untouched, "
                  f"{reason} — {why[reason]}: {', '.join(markers)}", file=sys.stderr)
    if unknown:
        print(f"[md_marker_convert] WARNING: unmapped tokens left untouched: "
              f"{sorted(unknown)}", file=sys.stderr)
    failing = bool(left.get("UNMAPPED") or left.get("MALFORMED") or unknown)
    if failing and not args.allow_partial:
        print(f"[md_marker_convert] FAIL: partial conversion written to {args.output}. Fix the map, "
              f"or pass --allow-partial to accept it.", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
